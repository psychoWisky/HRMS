"""Excel (.xlsx) export for reports.

The client requirement is explicit: reports download as Excel, never CSV.
"""
import io

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="14532D", end_color="14532D", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


_INVALID_SHEET_CHARS = str.maketrans({c: " " for c in "/\\?*[]:"})


def build_workbook(title: str, columns: list[str], rows: list[list]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    safe_title = (title or "Report").translate(_INVALID_SHEET_CHARS).strip()
    ws.title = safe_title[:31] or "Report"

    for col_index, header in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_index, value=value)

    widths = [len(str(c)) for c in columns]
    for row in rows:
        for i, value in enumerate(row):
            widths[i] = max(widths[i], len(str(value)) if value is not None else 0)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)

    ws.freeze_panes = "A2"
    return wb


def xlsx_response(title: str, columns: list[str], rows: list[list], filename: str) -> StreamingResponse:
    wb = build_workbook(title, columns, rows)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    safe_name = filename if filename.endswith(".xlsx") else f"{filename}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}"'},
    )
