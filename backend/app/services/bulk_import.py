"""Bulk-import current employees from an .xlsx spreadsheet.

For onboarding the many employees who already work at the university when
the HRMS goes live — as opposed to the public submission flow, which is for
future joiners. Proof documents for the new date-of-birth / date-of-joining
/ retirement fields are not part of the sheet; HR attaches those afterwards
per employee via the KYC document upload.
"""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.models import Designation, Employee, EmploymentStatus, Location, OrgUnit
from app.services.retirement import calc_retirement_date

COLUMNS = [
    "full_name",
    "hrms_employee_id",
    "gender",
    "date_of_birth",
    "official_email",
    "phone",
    "org_unit_name",
    "location_name",
    "designation_name",
    "pay_scale",
    "date_of_joining",
    "date_of_joining_aau_avfu",
    "date_of_joining_present_post",
    "expected_date_of_retirement",
    "remarks",
]

REQUIRED = {"full_name", "org_unit_name"}


def build_template() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(COLUMNS)
    ws.append(
        [
            "Dr. Jane Doe",
            "",
            "Female",
            "1985-06-15",
            "jane.doe@example.com",
            "9999999999",
            "College of Veterinary Science, Khanapara",
            "Khanapara, Guwahati",
            "Professor",
            "Level 14",
            "2010-08-01",
            "2010-08-01",
            "2022-01-01",
            "",
            "",
        ]
    )
    for i, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
            18, len(header) + 2
        )
    return wb


def _parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date '{text}' (use YYYY-MM-DD)")


class RowError(Exception):
    pass


def parse_rows(contents: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c else "" for c in next(rows_iter, [])]
    col_index = {name: i for i, name in enumerate(header) if name}

    missing = [c for c in REQUIRED if c not in col_index]
    if missing:
        raise RowError(f"Missing required column(s): {', '.join(missing)}")

    parsed = []
    for row in rows_iter:
        if row is None or all(v in (None, "") for v in row):
            continue
        record: dict = {}
        for col in COLUMNS:
            if col not in col_index:
                record[col] = None
                continue
            idx = col_index[col]
            record[col] = row[idx] if idx < len(row) else None
        parsed.append(record)
    return parsed


def create_employee_from_row(
    db: Session, row: dict, row_number: int, generate_id
) -> Employee:
    full_name = (row.get("full_name") or "").strip()
    if not full_name:
        raise RowError(f"Row {row_number}: full_name is required")

    org_unit_name = (row.get("org_unit_name") or "").strip()
    org_unit = (
        db.query(OrgUnit).filter(OrgUnit.name.ilike(org_unit_name)).first()
        if org_unit_name
        else None
    )
    if org_unit_name and org_unit is None:
        raise RowError(f"Row {row_number}: org unit '{org_unit_name}' not found")

    location = None
    location_name = (row.get("location_name") or "").strip()
    if location_name:
        location = db.query(Location).filter(Location.name.ilike(location_name)).first()
        if location is None:
            raise RowError(f"Row {row_number}: campus/location '{location_name}' not found")

    designation = None
    designation_name = (row.get("designation_name") or "").strip()
    if designation_name:
        designation = (
            db.query(Designation).filter(Designation.name.ilike(designation_name)).first()
        )
        if designation is None:
            raise RowError(f"Row {row_number}: designation '{designation_name}' not found")

    try:
        dob = _parse_date(row.get("date_of_birth"))
        doj = _parse_date(row.get("date_of_joining"))
        doj_aau_avfu = _parse_date(row.get("date_of_joining_aau_avfu"))
        doj_present_post = _parse_date(row.get("date_of_joining_present_post"))
        retirement = _parse_date(row.get("expected_date_of_retirement"))
    except ValueError as exc:
        raise RowError(f"Row {row_number}: {exc}") from exc

    if retirement is None:
        retirement = calc_retirement_date(dob, designation.rank_level if designation else None)

    hrms_id = (row.get("hrms_employee_id") or "").strip() or generate_id(
        db, org_unit.id if org_unit else None
    )
    if db.query(Employee).filter(Employee.hrms_employee_id == hrms_id).first():
        raise RowError(f"Row {row_number}: HRMS Employee ID '{hrms_id}' is already in use")

    employee = Employee(
        hrms_employee_id=hrms_id,
        full_name=full_name,
        gender=(row.get("gender") or "").strip(),
        date_of_birth=dob,
        date_of_joining_aau_avfu=doj_aau_avfu,
        date_of_joining_present_post=doj_present_post,
        expected_date_of_retirement=retirement,
        official_email=(row.get("official_email") or "").strip().lower(),
        phone=(row.get("phone") or "").strip(),
        org_unit_id=org_unit.id if org_unit else None,
        location_id=location.id if location else None,
        designation_id=designation.id if designation else None,
        pay_scale=(row.get("pay_scale") or "").strip(),
        date_of_joining=doj,
        employment_status=EmploymentStatus.active,
        remarks=(row.get("remarks") or "").strip(),
    )
    return employee
